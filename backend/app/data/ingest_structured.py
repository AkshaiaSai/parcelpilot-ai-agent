"""
Ingest structured data from ParcelPilot_Assessment_Data.xlsx into SQLite.

Reads the accounts, orders, and tickets sheets from the xlsx file,
creates corresponding SQLite tables, and inserts all data rows.
Also creates an action_log table for the two-phase action system.

Usage:
    python -m app.data.ingest_structured
"""

import sqlite3
import sys
from pathlib import Path

import openpyxl

# Allow running as module or standalone
try:
    from app.config import RAW_DATA_DIR, SQLITE_DB_PATH, STORAGE_DIR
except ImportError:
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))
    from app.config import RAW_DATA_DIR, SQLITE_DB_PATH, STORAGE_DIR


XLSX_FILE = RAW_DATA_DIR / "ParcelPilot_Assessment_Data.xlsx"

# --- Schema definitions ---

ACCOUNTS_SCHEMA = """
CREATE TABLE IF NOT EXISTS accounts (
    account_id TEXT PRIMARY KEY,
    account_name TEXT NOT NULL,
    plan TEXT NOT NULL,
    status TEXT NOT NULL,
    csm TEXT,
    contract_file TEXT,
    premium_support BOOLEAN DEFAULT FALSE,
    notes TEXT
)
"""

ORDERS_SCHEMA = """
CREATE TABLE IF NOT EXISTS orders (
    order_id TEXT PRIMARY KEY,
    account_id TEXT NOT NULL,
    carrier TEXT,
    status TEXT NOT NULL,
    booked_at TEXT,
    pickup_window_start TEXT,
    pickup_window_end TEXT,
    pickup_actual_at TEXT,
    shipment_fee_inr REAL,
    carrier_fault BOOLEAN DEFAULT FALSE,
    customer_fault BOOLEAN DEFAULT FALSE,
    cancellation_requested_at TEXT,
    notes TEXT,
    FOREIGN KEY (account_id) REFERENCES accounts(account_id)
)
"""

TICKETS_SCHEMA = """
CREATE TABLE IF NOT EXISTS tickets (
    ticket_id TEXT PRIMARY KEY,
    account_id TEXT NOT NULL,
    created_at TEXT,
    status TEXT NOT NULL,
    subject TEXT,
    description TEXT,
    channel TEXT,
    assigned_to TEXT,
    last_customer_message_at TEXT,
    historical_resolution TEXT,
    FOREIGN KEY (account_id) REFERENCES accounts(account_id)
)
"""

ACTION_LOG_SCHEMA = """
CREATE TABLE IF NOT EXISTS action_log (
    action_id TEXT PRIMARY KEY,
    action_type TEXT NOT NULL,
    status TEXT NOT NULL DEFAULT 'pending',
    details TEXT,
    created_by TEXT,
    created_at TEXT,
    confirmed_at TEXT,
    related_ticket_id TEXT,
    related_order_id TEXT,
    related_account_id TEXT
)
"""


def _read_sheet(wb: openpyxl.Workbook, sheet_name: str) -> list[dict]:
    """Read a sheet into a list of dicts, skipping empty rows."""
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []

    headers = [h for h in rows[0] if h is not None]
    num_cols = len(headers)
    data = []

    for row in rows[1:]:
        # Skip entirely empty rows
        values = row[:num_cols]
        if all(v is None for v in values):
            continue
        record = {}
        for col_idx, header in enumerate(headers):
            val = values[col_idx] if col_idx < len(values) else None
            # Convert datetime objects to ISO strings
            if hasattr(val, "isoformat"):
                val = val.isoformat()
            # Convert booleans to int for SQLite
            if isinstance(val, bool):
                val = int(val)
            record[header] = val
        data.append(record)

    return data


def _insert_rows(conn: sqlite3.Connection, table: str, rows: list[dict]):
    """Insert rows into a table using parameterized queries."""
    if not rows:
        return
    columns = list(rows[0].keys())
    placeholders = ", ".join(["?"] * len(columns))
    col_names = ", ".join(columns)
    sql = f"INSERT OR REPLACE INTO {table} ({col_names}) VALUES ({placeholders})"

    for row in rows:
        values = [row.get(col) for col in columns]
        conn.execute(sql, values)


def ingest():
    """Main ingestion function."""
    print(f"📂 Reading xlsx: {XLSX_FILE}")
    if not XLSX_FILE.exists():
        print(f"❌ File not found: {XLSX_FILE}")
        sys.exit(1)

    wb = openpyxl.load_workbook(str(XLSX_FILE), read_only=True)

    # Read data sheets
    accounts = _read_sheet(wb, "accounts")
    orders = _read_sheet(wb, "orders")
    tickets = _read_sheet(wb, "tickets")
    wb.close()

    print(f"  → accounts: {len(accounts)} rows")
    print(f"  → orders:   {len(orders)} rows")
    print(f"  → tickets:  {len(tickets)} rows")

    # Ensure storage directory exists
    STORAGE_DIR.mkdir(parents=True, exist_ok=True)

    # Remove existing DB to start fresh
    if SQLITE_DB_PATH.exists():
        SQLITE_DB_PATH.unlink()
        print(f"  🗑️  Removed existing DB")

    # Create DB and tables
    conn = sqlite3.connect(str(SQLITE_DB_PATH))
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")

    conn.executescript(ACCOUNTS_SCHEMA)
    conn.executescript(ORDERS_SCHEMA)
    conn.executescript(TICKETS_SCHEMA)
    conn.executescript(ACTION_LOG_SCHEMA)

    # Insert data
    _insert_rows(conn, "accounts", accounts)
    _insert_rows(conn, "orders", orders)
    _insert_rows(conn, "tickets", tickets)

    conn.commit()

    # Verify
    print(f"\n✅ SQLite DB created at: {SQLITE_DB_PATH}")
    print(f"   Size: {SQLITE_DB_PATH.stat().st_size:,} bytes\n")

    cursor = conn.cursor()
    for table in ["accounts", "orders", "tickets", "action_log"]:
        cursor.execute(f"SELECT COUNT(*) FROM {table}")
        count = cursor.fetchone()[0]
        print(f"   {table}: {count} rows")

    # Print sample data
    print("\n📊 Sample data:")
    for table in ["accounts", "orders", "tickets"]:
        print(f"\n   --- {table} ---")
        cursor.execute(f"SELECT * FROM {table} LIMIT 2")
        cols = [desc[0] for desc in cursor.description]
        print(f"   Columns: {cols}")
        for row in cursor.fetchall():
            print(f"   {dict(zip(cols, row))}")

    conn.close()
    print("\n✅ Structured data ingestion complete!")


if __name__ == "__main__":
    ingest()
